from __future__ import annotations

import csv
import hashlib
import io
import ipaddress
import json
import re
import socket
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from PIL import Image, ImageOps

try:
    from docx import Document
except ImportError:  # pragma: no cover - surfaced as a user-facing dependency error
    Document = None

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None

from openpyxl import load_workbook

from app.services.member_import_parser import parse_text_to_drafts, table_rows_to_drafts
from app.services.member_import_validator import validate_import_draft

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_PIXELS = 30_000_000
SUPPORTED_DOCUMENTS = {".docx", ".xlsx", ".xls", ".csv", ".txt"}
SUPPORTED_IMAGES = {".jpg", ".jpeg", ".png", ".webp"}

FIELD_ALIASES: dict[str, set[str]] = {
    "name": {"姓名", "会员姓名", "联系人", "名字", "name", "membername", "contact"},
    "organization_name": {"机构", "公司", "单位", "所在机构", "企业", "组织", "organization", "company", "employer"},
    "title": {"职位", "职务", "岗位", "角色", "头衔", "title", "position", "role"},
    "mobile": {"手机", "手机号", "电话", "联系电话", "联系方式", "mobile", "phone", "tel"},
    "email": {"邮箱", "电子邮箱", "email", "mail", "e-mail"},
    "wechat": {"微信", "微信号", "wechat", "weixin"},
    "city": {"城市", "地区", "所在地", "区域", "city", "region", "location"},
    "industry_tags": {"关注赛道", "行业标签", "赛道", "行业", "产业方向", "industry", "industrytags"},
    "expertise_tags": {"专业能力", "能力标签", "专长", "专业方向", "expertise", "skills", "ability"},
    "offered_resources": {"可提供资源", "资源", "供给", "优势资源", "offeredresources", "offering", "resources"},
    "cooperation_needs": {"合作需求", "当前需求", "需求", "寻求资源", "cooperationneeds", "needs", "demand"},
    "self_introduction": {"个人简介", "简介", "自我介绍", "背景", "bio", "introduction", "profile"},
    "referral_source": {"来源", "加入来源", "渠道", "referralsource", "source"},
    "referrer_name": {"推荐人", "引荐人", "referrer", "introducedby"},
    "preferred_contact_method": {"联系偏好", "希望联系方式", "preferredcontactmethod", "contactpreference"},
    "member_level": {"会员等级", "等级", "memberlevel", "level"},
    "member_status": {"会员状态", "状态", "memberstatus", "status"},
    "owner": {"负责人", "运营负责人", "owner", "manager"},
}

CANONICAL_FIELDS = list(FIELD_ALIASES)
LABEL_LOOKUP = {
    "name": "姓名",
    "organization_name": "机构",
    "title": "职位",
    "mobile": "手机",
    "email": "邮箱",
    "wechat": "微信",
    "city": "城市",
    "industry_tags": "关注赛道",
    "expertise_tags": "专业能力",
    "offered_resources": "可提供资源",
    "cooperation_needs": "合作需求",
    "self_introduction": "个人简介",
    "referral_source": "来源",
    "referrer_name": "推荐人",
    "preferred_contact_method": "联系偏好",
    "member_level": "会员等级",
    "member_status": "会员状态",
    "owner": "负责人",
}

MOBILE_RE = re.compile(r"(?<!\d)(?:\+?86[-\s]?)?1[3-9]\d{9}(?!\d)")
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
KEY_VALUE_RE = re.compile(r"^\s*([^:：]{1,24})\s*[:：]\s*(.*?)\s*$")


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u3000", " ").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    return re.sub(r"[\s_\-—–:：/\\（）()\[\]【】]+", "", text)


def canonical_field(header: Any) -> str | None:
    normalized = normalize_header(header)
    if not normalized:
        return None
    for field, aliases in FIELD_ALIASES.items():
        if normalized in {normalize_header(item) for item in aliases}:
            return field
    return None


def default_draft() -> dict[str, Any]:
    draft = {field: "" for field in CANONICAL_FIELDS}
    draft.update(
        {
            "member_level": "standard",
            "member_status": "active",
            "source_locator": "",
            "source_text": "",
            "warnings": [],
            "image_bytes": None,
            "image_name": "",
        }
    )
    return draft


def _table_to_drafts(rows: list[list[Any]], *, locator_prefix: str) -> list[dict[str, Any]]:
    rows = [[clean_text(cell) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    header_index = 0
    best_mapping: dict[int, str] = {}
    for idx, row in enumerate(rows[:10]):
        mapping = {col: field for col, value in enumerate(row) if (field := canonical_field(value))}
        if len(mapping) > len(best_mapping):
            header_index, best_mapping = idx, mapping
    drafts: list[dict[str, Any]] = []
    if len(best_mapping) >= 2:
        for row_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            draft = default_draft()
            for col, field in best_mapping.items():
                if col < len(row):
                    draft[field] = row[col]
            if not any(draft.get(field) for field in ("name", "organization_name", "mobile", "email", "self_introduction")):
                continue
            draft["source_locator"] = f"{locator_prefix}第{row_no}行"
            draft["source_text"] = " | ".join(cell for cell in row if cell)
            drafts.append(draft)
        return drafts

    # No reliable header: interpret each non-empty row as one person using conservative positional fields.
    for row_no, row in enumerate(rows, start=1):
        values = [cell for cell in row if cell]
        if not values:
            continue
        draft = default_draft()
        draft["name"] = values[0]
        if len(values) > 1:
            draft["organization_name"] = values[1]
        if len(values) > 2:
            draft["title"] = values[2]
        remainder = "；".join(values[3:])
        if remainder:
            draft["self_introduction"] = remainder
        _extract_contacts(draft, " ".join(values))
        draft["warnings"].append("未识别到明确表头，已按姓名、机构、职位的顺序生成草稿，请逐条核对。")
        draft["source_locator"] = f"{locator_prefix}第{row_no}行"
        draft["source_text"] = " | ".join(values)
        drafts.append(draft)
    return drafts


def _extract_contacts(draft: dict[str, Any], text: str) -> None:
    if not draft.get("mobile"):
        mobile = MOBILE_RE.search(text)
        if mobile:
            draft["mobile"] = re.sub(r"\D", "", mobile.group())[-11:]
    if not draft.get("email"):
        email = EMAIL_RE.search(text)
        if email:
            draft["email"] = email.group()


def _parse_key_value_block(block: str, locator: str) -> dict[str, Any] | None:
    draft = default_draft()
    unmatched: list[str] = []
    matched = 0
    for line in block.splitlines():
        line = line.strip()
        if not line:
            continue
        pair = KEY_VALUE_RE.match(line)
        if pair and (field := canonical_field(pair.group(1))):
            draft[field] = pair.group(2).strip()
            matched += 1
        else:
            unmatched.append(line)
    if matched == 0:
        return None
    if unmatched and not draft["self_introduction"]:
        draft["self_introduction"] = "\n".join(unmatched)
    draft["source_locator"] = locator
    draft["source_text"] = block
    _extract_contacts(draft, block)
    return draft


def parse_text(text: str, *, locator_prefix: str = "粘贴文本") -> list[dict[str, Any]]:
    text = clean_text(text)
    if not text:
        return []

    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) >= 2 and any("\t" in line for line in lines[:3]):
        rows = [line.split("\t") for line in lines]
        result = _table_to_drafts(rows, locator_prefix=locator_prefix)
        if result:
            return result

    # CSV-like pasted content.
    if len(lines) >= 2 and any("," in line or "，" in line for line in lines[:3]):
        normalized = text.replace("，", ",")
        try:
            rows = list(csv.reader(io.StringIO(normalized)))
            result = _table_to_drafts(rows, locator_prefix=locator_prefix)
            if result and len(result) > 1:
                return result
        except csv.Error:
            pass

    blocks = [part.strip() for part in re.split(r"\n\s*\n+", text) if part.strip()]
    drafts: list[dict[str, Any]] = []
    for idx, block in enumerate(blocks, start=1):
        parsed = _parse_key_value_block(block, f"{locator_prefix}第{idx}段")
        if parsed:
            drafts.append(parsed)
    if drafts:
        return drafts

    # A list where each line is a compact member record.
    if len(lines) > 1 and all(len(line) < 220 for line in lines):
        compact: list[dict[str, Any]] = []
        for idx, line in enumerate(lines, start=1):
            values = [v.strip() for v in re.split(r"[\t|；;]+", line) if v.strip()]
            if len(values) >= 2:
                draft = default_draft()
                draft["name"] = values[0]
                draft["organization_name"] = values[1] if len(values) > 1 else ""
                draft["title"] = values[2] if len(values) > 2 else ""
                if len(values) > 3:
                    draft["self_introduction"] = "；".join(values[3:])
                _extract_contacts(draft, line)
                draft["source_locator"] = f"{locator_prefix}第{idx}行"
                draft["source_text"] = line
                draft["warnings"].append("紧凑名单按姓名、机构、职位顺序解析，请核对字段归属。")
                compact.append(draft)
        if compact:
            return compact

    draft = default_draft()
    draft["source_locator"] = locator_prefix
    draft["source_text"] = text
    draft["self_introduction"] = text
    _extract_contacts(draft, text)
    first_line = lines[0].strip() if lines else ""
    if len(first_line) <= 30 and not re.search(r"[，。；,:：]", first_line):
        draft["name"] = first_line
        draft["self_introduction"] = "\n".join(lines[1:])
    draft["warnings"].append("文本结构不足，未自动猜测复杂字段，请补充姓名、机构和职位。")
    return [draft]


def _table_to_drafts(rows: list[list[Any]], *, locator_prefix: str) -> list[dict[str, Any]]:
    """v0.5C parser: never map no-header columns by position."""
    return table_rows_to_drafts(rows, locator_prefix=locator_prefix)


def parse_text(text: str, *, locator_prefix: str = "粘贴文本") -> list[dict[str, Any]]:
    """v0.5C parser: extract only labeled or semantically explicit fields."""
    return parse_text_to_drafts(text, locator_prefix=locator_prefix)


def parse_csv_file(path: Path) -> list[dict[str, Any]]:
    raw = path.read_bytes()
    text = ""
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV。")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _table_to_drafts(rows, locator_prefix="CSV ")


def parse_xlsx_file(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(path, data_only=True)
    drafts: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    row_to_global: dict[tuple[str, int], int] = {}
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        parsed = _table_to_drafts(rows, locator_prefix=f"工作表 {sheet.title} ")
        base = len(drafts)
        drafts.extend(parsed)
        # Source locator created by table parser contains the original row number.
        for offset, draft in enumerate(parsed):
            match = re.search(r"第(\d+)行", draft["source_locator"])
            if match:
                row_to_global[(sheet.title, int(match.group(1)))] = base + offset
        for image in getattr(sheet, "_images", []):
            try:
                anchor_row = int(image.anchor._from.row) + 1
                data = image._data()
                ext = (getattr(image, "format", "png") or "png").lower()
                target = row_to_global.get((sheet.title, anchor_row))
                images.append(
                    {
                        "bytes": data,
                        "filename": f"{sheet.title}_row{anchor_row}.{ext}",
                        "draft_index": target,
                        "source_locator": f"工作表 {sheet.title} 第{anchor_row}行嵌入图片",
                    }
                )
            except Exception:
                continue
    return drafts, images


def parse_xls_file(path: Path) -> list[dict[str, Any]]:
    if xlrd is None:
        raise ValueError("缺少 xlrd 依赖，请重新运行 setup_windows.bat 后再导入 .xls。")
    workbook = xlrd.open_workbook(path)
    drafts: list[dict[str, Any]] = []
    for sheet in workbook.sheets():
        rows = [sheet.row_values(row_idx) for row_idx in range(sheet.nrows)]
        drafts.extend(_table_to_drafts(rows, locator_prefix=f"工作表 {sheet.name} "))
    return drafts


def parse_docx_file(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if Document is None:
        raise ValueError("缺少 python-docx 依赖，请重新运行 setup_windows.bat。")
    document = Document(path)
    drafts: list[dict[str, Any]] = []
    for table_idx, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        drafts.extend(_table_to_drafts(rows, locator_prefix=f"Word 表格{table_idx} "))

    paragraphs = [clean_text(p.text) for p in document.paragraphs if clean_text(p.text)]
    if paragraphs:
        paragraph_text = "\n".join(paragraphs)
        parsed = parse_text(paragraph_text, locator_prefix="Word 正文")
        # Avoid duplicating the same table content echoed into paragraphs.
        existing_names = {normalize_header(d.get("name")) for d in drafts if d.get("name")}
        for draft in parsed:
            if draft.get("name") and normalize_header(draft["name"]) in existing_names:
                continue
            drafts.append(draft)

    images: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
        for idx, name in enumerate(names):
            data = archive.read(name)
            images.append(
                {
                    "bytes": data,
                    "filename": Path(name).name,
                    "draft_index": idx if idx < len(drafts) else None,
                    "source_locator": f"Word 内嵌图片 {idx + 1}",
                }
            )
    if images and len(images) != len(drafts):
        for draft in drafts:
            draft["warnings"].append("Word 中照片数量与人物数量不一致，头像关联需要人工确认。")
    return drafts, images


def parse_uploaded_document(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return parse_csv_file(path), []
    if suffix == ".txt":
        raw = path.read_bytes()
        text = ""
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            raise ValueError("文本文件编码无法识别，请另存为 UTF-8。")
        return parse_text(text, locator_prefix="文本文件"), []
    if suffix == ".xlsx":
        return parse_xlsx_file(path)
    if suffix == ".xls":
        return parse_xls_file(path), []
    if suffix == ".docx":
        return parse_docx_file(path)
    if suffix == ".doc":
        raise ValueError("旧版 .doc 暂不支持，请在 Word 中另存为 .docx 后导入。")
    raise ValueError("不支持的文件格式。支持 .docx、.xlsx、.xls、.csv、.txt。")


def validate_draft(draft: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    name = clean_text(draft.get("name"))
    if not name:
        warnings.append("姓名不能为空")
    if len(name) > 80:
        warnings.append("姓名过长，可能混入职位或简介")
    if any(word in name for word in ("管理团队", "核心团队", "董事会", "关于我们", "联系我们", "新闻中心", "团队介绍")):
        warnings.append("姓名疑似栏目名称，不能直接保存")
    if re.search(r"[；;、]", name):
        warnings.append("姓名中包含多个分隔符，疑似多人")
    mobile = re.sub(r"\D", "", clean_text(draft.get("mobile")))
    if mobile and len(mobile) not in {11, 12, 13, 14, 15}:
        warnings.append("手机号格式需要核对")
    email = clean_text(draft.get("email"))
    if email and not EMAIL_RE.fullmatch(email):
        warnings.append("邮箱格式需要核对")
    if len(clean_text(draft.get("organization_name"))) > 180:
        warnings.append("机构字段过长，可能混入多家机构或履历")
    return warnings


def validate_draft(draft: dict[str, Any]) -> list[str]:
    """v0.5C quality gate with conservative name/org/title checks."""
    return validate_import_draft(draft)


def _host_is_public(hostname: str) -> bool:
    if not hostname or hostname.lower() in {"localhost", "localhost.localdomain"}:
        return False
    try:
        addresses = socket.getaddrinfo(hostname, None)
    except socket.gaierror:
        return False
    for info in addresses:
        address = ipaddress.ip_address(info[4][0])
        if address.is_private or address.is_loopback or address.is_link_local or address.is_reserved or address.is_multicast:
            return False
    return True


def validate_public_url(url: str) -> str:
    url = clean_text(url)
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("请输入有效的公开 http/https 网页地址。")
    if not _host_is_public(parsed.hostname):
        raise ValueError("出于安全原因，不能读取本机或内网地址。")
    return url




def _safe_http_get(url: str, *, timeout: float, max_bytes: int) -> httpx.Response:
    headers = {"User-Agent": "Mozilla/5.0 (compatible; Q-BAY Member Import/1.0)"}
    current = validate_public_url(url)
    with httpx.Client(timeout=timeout, follow_redirects=False, headers=headers) as client:
        for _ in range(6):
            response = client.get(current)
            if response.status_code in {301, 302, 303, 307, 308}:
                location = response.headers.get("location")
                if not location:
                    response.raise_for_status()
                current = validate_public_url(urljoin(current, location))
                continue
            response.raise_for_status()
            length = response.headers.get("content-length")
            if length and length.isdigit() and int(length) > max_bytes:
                raise ValueError("远程内容超过允许大小。")
            if len(response.content) > max_bytes:
                raise ValueError("远程内容超过允许大小。")
            validate_public_url(str(response.url))
            return response
    raise ValueError("网页重定向次数过多。")

def fetch_web_image_candidates(url: str, *, limit: int = 30) -> tuple[str, list[dict[str, str]]]:
    url = validate_public_url(url)
    response = _safe_http_get(url, timeout=15.0, max_bytes=5 * 1024 * 1024)
    content_type = response.headers.get("content-type", "")
    if "html" not in content_type.lower():
        raise ValueError("该地址不是可读取的 HTML 网页。")
    soup = BeautifulSoup(response.text, "html.parser")
    title = clean_text(soup.title.string if soup.title and soup.title.string else "")
    candidates: list[dict[str, str]] = []
    seen: set[str] = set()

    def add(image_url: str, alt: str = "", source: str = "img") -> None:
        absolute = urljoin(str(response.url), image_url.strip())
        parsed = urlparse(absolute)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            return
        if not _host_is_public(parsed.hostname):
            return
        if absolute in seen or absolute.lower().endswith(".svg"):
            return
        seen.add(absolute)
        candidates.append({"image_url": absolute, "alt_text": clean_text(alt)[:300], "source": source})

    for selector, attr in [
        ('meta[property="og:image"]', "content"),
        ('meta[name="twitter:image"]', "content"),
        ('meta[property="twitter:image"]', "content"),
    ]:
        for node in soup.select(selector):
            if node.get(attr):
                add(node.get(attr), title, "metadata")
    for image in soup.find_all("img"):
        source = image.get("src") or image.get("data-src") or image.get("data-original")
        if not source:
            continue
        width = str(image.get("width") or "")
        height = str(image.get("height") or "")
        if width.isdigit() and height.isdigit() and (int(width) < 80 or int(height) < 80):
            continue
        add(source, image.get("alt") or image.get("title") or "", "img")
        if len(candidates) >= limit:
            break
    return title, candidates[:limit]


def normalize_and_store_image(
    data: bytes,
    *,
    storage_root: Path,
    preferred_name: str = "avatar",
) -> dict[str, Any]:
    if not data or len(data) > MAX_IMAGE_BYTES:
        raise ValueError("图片为空或超过 10MB。")
    digest = hashlib.sha256(data).hexdigest()
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        image = Image.open(io.BytesIO(data))
        image.verify()
        image = Image.open(io.BytesIO(data))
        image = ImageOps.exif_transpose(image)
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError("图片像素过大。")
        if image.mode not in {"RGB", "RGBA"}:
            image = image.convert("RGBA" if "transparency" in image.info else "RGB")
    except Exception as exc:
        raise ValueError("无法识别为有效的 JPG、PNG 或 WEBP 图片。") from exc

    stamp = datetime.now()
    relative_dir = Path(f"{stamp:%Y}") / f"{stamp:%m}"
    target_dir = storage_root / relative_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{digest[:16]}_{re.sub(r'[^A-Za-z0-9_-]+', '_', Path(preferred_name).stem)[:30] or 'avatar'}"
    stored_relative = relative_dir / f"{stem}.webp"
    thumb_relative = relative_dir / f"{stem}_thumb.webp"
    target = storage_root / stored_relative
    thumb = storage_root / thumb_relative
    if not target.exists():
        image.save(target, "WEBP", quality=90, method=6)
        thumbnail = image.copy()
        thumbnail.thumbnail((360, 360), Image.Resampling.LANCZOS)
        thumbnail.save(thumb, "WEBP", quality=86, method=6)
    return {
        "stored_filename": stored_relative.as_posix(),
        "thumbnail_filename": thumb_relative.as_posix(),
        "mime_type": "image/webp",
        "size_bytes": target.stat().st_size,
        "width": image.width,
        "height": image.height,
        "sha256": digest,
    }


def download_public_image(url: str) -> tuple[bytes, str, str]:
    url = validate_public_url(url)
    response = _safe_http_get(url, timeout=20.0, max_bytes=MAX_IMAGE_BYTES)
    content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
    if not content_type.startswith("image/"):
        raise ValueError("候选地址不是图片。")
    filename = Path(urlparse(str(response.url)).path).name or "web-avatar"
    return response.content, filename, content_type


def field_labels() -> dict[str, str]:
    return dict(LABEL_LOOKUP)
