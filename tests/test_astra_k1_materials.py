"""Direct K1 checks; explicit isolated seed and no outbound network."""
import importlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from contextlib import closing
from unittest.mock import patch

from sqlalchemy import create_engine,event,text
from sqlalchemy.orm import Session
from fastapi import HTTPException

from app.services.knowledge_service import KnowledgeService,policy_view,process_material_task
from app.services.knowledge_material_parser import parse_material
from scripts.migrate_db import _apply_migration,MIGRATIONS


class MaterialsTest(unittest.TestCase):
    def setUp(self):
        self.authentication=patch.dict(os.environ,{'APP_AUTH_DISABLED':'false'});self.authentication.start();self.addCleanup(self.authentication.stop)
        seed=Path(os.environ['K1_TEST_SEED']).resolve()
        assert seed.name!='app.db' and 'acceptance' in seed.parts
        self.tmp=tempfile.TemporaryDirectory(prefix='k1_',dir=seed.parent)
        self.path=Path(self.tmp.name)/'test.db'
        with closing(sqlite3.connect(seed.as_uri()+'?mode=ro',uri=True)) as src,closing(sqlite3.connect(self.path)) as dst:
            src.backup(dst)
            dst.execute('pragma foreign_keys=on');dst.execute('begin immediate')
            _apply_migration(dst,MIGRATIONS[-1]);dst.commit()
            self.uid=dst.execute("select id from v05a_users where role='admin' order by id limit 1").fetchone()[0]
        self.engine=create_engine('sqlite:///'+self.path.as_posix())
        @event.listens_for(self.engine,'connect')
        def fk(conn,_):conn.execute('pragma foreign_keys=on')
        self.db=Session(self.engine);self.s=KnowledgeService(self.db,user_id=self.uid)
        self.network=patch('socket.create_connection',side_effect=AssertionError('test_network_forbidden'));self.network.start()

    def tearDown(self):
        self.network.stop();self.db.close();self.engine.dispose();self.tmp.cleanup()

    def upload(self,body='第一条 适用条件\n仅适用于依法登记的研发机构。\n第二条 支持标准\n符合第一条条件，最高不超过实际投入的百分之十。',mid=None):
        fields={'title':'K1隔离版本fixture，不是真实政策','category':'政策/研发资助','access_scope':'PRIVATE'}
        result=self.s.register_material(fields,body.encode(),'fixture.txt',mid);self.db.commit()
        process_material_task({'material_id':result,'owner_user_id':self.uid},self.path)
        return result

    def confirm(self,c,action='draft',**kw):
        result=self.s.review_candidate(c['id'],{'edit_revision':c['edit_revision'],'acknowledge':'1',**kw},action)
        self.db.commit();return result

    def test_complete_versions_idempotence_privacy(self):
        mid=self.upload();cs=self.s.candidates(mid);self.assertEqual(len(cs),2)
        first=self.confirm(cs[0],'publish');kid=first['knowledge_id']
        self.confirm(cs[0],'publish')
        original=self.s.get(kid);self.assertEqual(original['content_revision'],1)
        self.s.progress(kid,self.uid,'MASTERED','private fixture note');self.db.commit()
        self.upload(mid=mid)
        self.assertEqual(len(self.s.candidates(mid)),2)
        self.upload('第一条 适用条件\n仅适用于依法登记的研发机构。\n第二条 支持标准\n符合第一条条件，最高不超过实际投入的百分之二十。',mid)
        newer=next(c for c in self.s.candidates(mid) if not c['stale'] and c['target_knowledge_id']==kid)
        self.assertEqual(self.s.get(kid)['body'],original['body'])
        self.confirm(newer,'publish');self.assertEqual(self.s.get(kid)['content_revision'],2)
        history=self.s.rows('select content_json from knowledge_versions where knowledge_id=:k order by revision',k=kid)
        self.assertEqual(json.loads(history[0]['content_json'])['body'],original['body'])
        p=self.s.rows('select * from user_knowledge_progress where knowledge_id=:k',k=kid)[0]
        self.assertEqual(p['note'],'private fixture note');self.assertEqual(p['knowledge_revision'],1)
        stranger=KnowledgeService(self.db,user_id=-1)
        with self.assertRaises(HTTPException):stranger.get(kid,manager=True)
        self.assertNotIn(kid,[i['id'] for i in stranger.listing(manager=True)])
        with self.assertRaises(HTTPException):stranger.material(mid)
        with self.assertRaises(Exception):self.db.execute(text('update knowledge_versions set content_json=content_json'))
        self.db.rollback()

    def test_old_preview_and_target_conflict(self):
        mid=self.upload();c=self.s.candidates(mid)[0]
        self.upload('第一条 新版条件\n只适用于明确登记并经审核的主体。',mid)
        with self.assertRaises(HTTPException) as error:self.confirm(c)
        self.assertEqual(error.exception.status_code,409);self.db.rollback()
        c=next(c for c in self.s.candidates(mid) if not c['stale'])
        result=self.confirm(c,'publish');kid=result['knowledge_id']
        self.upload('第一条 新版条件\n明确登记并经审核，且完成备案。',mid)
        new=next(c for c in self.s.candidates(mid) if not c['stale'])
        item=self.s.get(kid);self.s.save({**item,'body':item['body']+'\n管理员的独立修正'},kid);self.db.commit()
        with self.assertRaises(HTTPException):self.confirm(new,'publish')
        self.db.rollback()

    def test_formats_bounds_and_policy_dates(self):
        from io import BytesIO
        from docx import Document
        doc=Document();doc.add_heading('第一章 隔离格式fixture',level=1);doc.add_paragraph('本资料仅用于解析验证，不是真实制度。')
        table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='适用条件';table.cell(0,1).text='支持口径';table.cell(1,0).text='条件A';table.cell(1,1).text='最高10%'
        stream=BytesIO();doc.save(stream)
        parsed=parse_material(stream.getvalue(),'fixture.docx')
        self.assertIn('条件A',str(parsed));self.assertIn('最高10%',str(parsed))
        self.assertIn('非页码',str(parsed));self.assertNotIn('PDF实际页序',str(parsed))
        from pypdf import PdfWriter
        pdf=PdfWriter();pdf.add_blank_page(width=200,height=200);stream=BytesIO();pdf.write(stream)
        with self.assertRaises(HTTPException):parse_material(stream.getvalue(),'scan.pdf')
        with self.assertRaises(HTTPException):parse_material(b'x'*(8*1024*1024+1),'large.txt')
        p=policy_view({'province':'上海市','city':'上海市','publication_date':'2000-01-01'})
        self.assertEqual(p['region_label'],'上海市');self.assertEqual(p['effect_status'],'UNVERIFIED');self.assertEqual(p['windows'],[])

    def test_migration_and_formal_exception_unchanged(self):
        self.db.rollback()
        with closing(sqlite3.connect(self.path)) as c:
            before=c.execute('pragma foreign_key_check').fetchall()
            _apply_migration(c,MIGRATIONS[-1]);c.commit()
            self.assertEqual(c.execute('pragma foreign_key_check').fetchall(),before)
            self.assertEqual(c.execute('pragma integrity_check').fetchone()[0],'ok')
            self.assertEqual(c.execute('select count(*) from knowledge_materials').fetchone()[0],0)

    def test_real_scheduler_tick_pause_and_lease(self):
        import time
        from app.services import collection_scheduler as scheduler
        from app.services.tasks import create_task
        mid=self.s.register_material({'title':'定时隔离fixture','category':'政策/招商','access_scope':'PRIVATE'},'第一条 规则\n仅隔离调度测试。'.encode(),'fixture.txt')
        self.db.execute(text('UPDATE v04g_monitoring_sources SET is_enabled=0'));self.db.commit()
        create_task('knowledge_material',queue_name='knowledge',payload={'material_id':mid,'owner_user_id':self.uid},idempotency_key='k1-test',db_path=self.path)
        try:
            self.assertTrue(scheduler.start_scheduler(force=True,db_path=self.path))
            scheduler._scheduler.reschedule_job('knowledge_materials',trigger='interval',seconds=.2)
            deadline=time.monotonic()+8
            while time.monotonic()<deadline:
                self.db.rollback()
                if self.s.candidates(mid):break
                time.sleep(.05)
            self.assertEqual(len(self.s.candidates(mid)),1)
            with patch.object(scheduler,'_instance_token','second-worker'):
                self.assertFalse(scheduler._heartbeat(self.path,claim=True))
            scheduler.set_automatic_collection(False,self.path)
            self.assertEqual(scheduler.runtime_status(self.path)['state'],'已暂停')
        finally:scheduler.stop_scheduler()
        self.assertFalse(scheduler.start_scheduler(db_path=self.path))
        self.assertEqual(scheduler.runtime_status(self.path)['state'],'已暂停')

    def test_merge_split_keeps_evidence(self):
        mid=self.upload();rows=sorted(self.s.candidates(mid),key=lambda c:c['id'])
        self.s.reshape_candidates(mid,[c['id'] for c in rows],{c['id']:c['edit_revision'] for c in rows});self.db.commit()
        merged=next(c for c in self.s.candidates(mid) if c['state']=='PENDING')
        self.assertEqual(len(merged['mappings']),sum(len(c['mappings']) for c in rows))
        self.s.reshape_candidates(mid,[merged['id']],{merged['id']:merged['edit_revision']},split_before='第二条');self.db.commit()
        split=[c for c in self.s.candidates(mid) if c['state']=='PENDING']
        self.assertEqual(len(split),2)
        self.assertTrue(all(c['mappings']==merged['mappings'] for c in split))

    def test_source_reuses_snapshot_without_network_or_intelligence(self):
        snapshot=self.s.rows("SELECT sn.id,sn.monitoring_source_id FROM v04g_source_snapshots sn JOIN v04g_monitoring_sources s ON s.id=sn.monitoring_source_id WHERE s.deactivated_at IS NULL AND sn.http_status=200 AND (sn.raw_html<>'' OR sn.cleaned_text<>'' OR sn.raw_content<>'') ORDER BY sn.id DESC LIMIT 1")[0]
        sid=snapshot['monitoring_source_id']
        self.db.execute(text("UPDATE v04g_monitoring_sources SET url='https://example.org/collection' WHERE id=:id"),{'id':sid})
        self.db.commit()  # Even an isolated copy retains immutable evidence unchanged.
        with patch('socket.getaddrinfo',return_value=[(2,1,6,'',('93.184.216.34',443))]):
            mid=self.s.register_material({'source_id':sid,'category':'政策/招商','access_scope':'PRIVATE'})
            config=json.loads(self.s.material(mid)['config_json']);config['snapshot_cursor']=snapshot['id']-1
            self.db.execute(text('UPDATE knowledge_materials SET config_json=:c WHERE id=:id'),{'c':json.dumps(config),'id':mid});self.db.commit()
            result=process_material_task({'material_id':mid,'owner_user_id':self.uid},self.path)
            self.assertEqual(result['created'],1)
            again=process_material_task({'material_id':mid,'owner_user_id':self.uid},self.path)
            self.assertEqual(again['created'],0)
        child=next(m for m in self.s.materials() if m['id']!=mid)
        self.assertEqual(json.loads(child['config_json'])['collection_snapshot_id'],snapshot['id'])
        process_material_task({'material_id':child['id'],'owner_user_id':self.uid},self.path)
        self.assertTrue(self.s.candidates(child['id']))

    def test_notes_and_explicit_target_reconfirmation(self):
        mid=self.upload();c=self.s.candidates(mid)[0];kid=self.confirm(c,'publish')['knowledge_id']
        self.s.annotate(kid,self.uid,{'content':'私有笔记不作为来源','type':'NOTE'});self.db.commit()
        self.upload('第一条 另一个结构\n全新资料内容。',mid)
        candidate=next(x for x in self.s.candidates(mid) if not x['stale'])
        with self.assertRaises(HTTPException):self.confirm(candidate,'publish',target_selection=f'{kid}:1')
        self.db.rollback()
        self.confirm(candidate,'edit',target_selection=f'{kid}:1')
        fresh=next(x for x in self.s.candidates(mid) if x['id']==candidate['id'])
        self.confirm(fresh,'publish',target_selection=f'{kid}:1')
        note=self.s.annotations(kid,self.uid)[0]
        self.assertEqual(note['content'],'私有笔记不作为来源');self.assertEqual(note['knowledge_revision'],1)

    def test_unknown_and_expired_heartbeat_do_not_claim_running(self):
        from app.services import collection_scheduler as scheduler
        before=scheduler.runtime_status(self.path);self.assertNotEqual(before['state'],'运行中')
        self.db.execute(text("INSERT OR REPLACE INTO worker_heartbeats(worker_id,queue_name,task_type,pid,hostname,status,started_at,heartbeat_at,metadata_json) VALUES ('collection-scheduler','knowledge','scheduler',1,'fixture','online','2000-01-01','2000-01-01',:m)"),{'m':json.dumps({'enabled':True,'jobs':[]})});self.db.commit()
        self.assertEqual(scheduler.runtime_status(self.path)['state'],'状态未知')
        self.assertEqual(self.s.rows("SELECT heartbeat_at FROM worker_heartbeats WHERE worker_id='collection-scheduler'")[0]['heartbeat_at'],'2000-01-01')

    def test_removed_section_remains_visible_without_deleting_knowledge(self):
        mid=self.upload();candidates=self.s.candidates(mid)
        for c in candidates:self.confirm(c,'publish')
        count=self.s.rows('SELECT COUNT(*) AS n FROM knowledge_items')[0]['n']
        self.upload('第一条 适用条件\n仅适用于依法登记的研发机构。',mid)
        config=json.loads(self.s.material(mid)['config_json'])
        self.assertIn('第二条 支持标准',config['last_change']['removed'])
        self.assertEqual(self.s.rows('SELECT COUNT(*) AS n FROM knowledge_items')[0]['n'],count)

    def test_development_bypass_never_exposes_private_material(self):
        mid=self.upload();kid=self.confirm(self.s.candidates(mid)[0],'publish')['knowledge_id']
        with patch.dict(os.environ,{'APP_AUTH_DISABLED':'true'}):
            with self.assertRaises(HTTPException):self.s.material(mid)
            with self.assertRaises(HTTPException):self.s.get(kid,manager=True)
            self.assertNotIn(mid,[m['id'] for m in self.s.materials()])
            self.assertNotIn(kid,[m['id'] for m in self.s.listing(manager=True)])
            with self.assertRaises(HTTPException):self.s.register_material({'title':'内部资料'},b'private','private.txt')


if __name__=='__main__':unittest.main()
