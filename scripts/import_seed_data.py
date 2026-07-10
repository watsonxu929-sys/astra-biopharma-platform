from __future__ import annotations
import json, sys
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.database import Base, SessionLocal, engine
from app.models import ActionItem, HistoricalEvent, ImportLog, Organization, Person, ProjectPool, Relation, Resource
DEFAULT_FILE = ROOT / "reference_documents" / "04_第二轮校准种子数据库_V1.1.xlsx"
def text(v):
    if v is None: return None
    v=str(v).strip()
    return v or None
def rows_as_dicts(ws):
    rows=list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers=[str(x).strip() if x is not None else "" for x in rows[0]]
    out=[]
    for n,row in enumerate(rows[1:],start=2):
        if all(v is None or str(v).strip()=="" for v in row): continue
        d={headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        d['_row_number']=n
        out.append(d)
    return out
def exists(db, model, external_id):
    return db.scalar(select(model).where(model.external_id==external_id)) is not None
def import_sheet(db, ws, name, model, mapper):
    r={'added':0,'skipped':0,'failed':0,'errors':[]}
    for row in rows_as_dicts(ws):
        try:
            obj=mapper(row)
            if not obj.external_id: raise ValueError('缺少ID')
            if exists(db,model,obj.external_id):
                r['skipped']+=1; continue
            db.add(obj); db.flush(); r['added']+=1
        except Exception as exc:
            db.rollback(); r['failed']+=1
            r['errors'].append(f"{name} 第{row.get('_row_number')}行：{exc}")
    db.commit(); return r
def main():
    fp=Path(sys.argv[1]) if len(sys.argv)>1 else DEFAULT_FILE
    if not fp.exists():
        print('ERROR: 找不到Excel文件：',fp)
        raise SystemExit(1)
    Base.metadata.create_all(bind=engine)
    wb=load_workbook(fp,data_only=True)
    required=['主体种子库','人物种子库','资源种子库','历史事件种子','项目池种子','关系种子库','校准行动清单']
    missing=[x for x in required if x not in wb.sheetnames]
    if missing:
        print('ERROR: 缺少工作表：','、'.join(missing)); raise SystemExit(1)
    db=SessionLocal(); summary={}
    try:
        summary['主体种子库']=import_sheet(db,wb['主体种子库'],'主体种子库',Organization,lambda r:Organization(external_id=text(r.get('主体ID')),standard_name=text(r.get('标准名称')) or '未命名主体',org_type=text(r.get('主体类型')),region=text(r.get('地域')),industry_tags=text(r.get('产业标签')),resources=text(r.get('可提供资源')),needs=text(r.get('潜在需求')),relationship_source=text(r.get('关系来源')),visibility=text(r.get('权限')) or '内部',verification_status=text(r.get('核验状态')) or '待核验'))
        summary['人物种子库']=import_sheet(db,wb['人物种子库'],'人物种子库',Person,lambda r:Person(external_id=text(r.get('人物ID')),name=text(r.get('姓名/群体')) or '未命名人物',public_role=text(r.get('公开角色')),organization_network=text(r.get('当前机构/网络')),ability_tags=text(r.get('能力标签')),value_provided=text(r.get('可提供价值')),relationship_source=text(r.get('关系来源')),visibility=text(r.get('权限')) or '内部',verification_status=text(r.get('核验状态')) or '待核验'))
        summary['资源种子库']=import_sheet(db,wb['资源种子库'],'资源种子库',Resource,lambda r:Resource(external_id=text(r.get('资源ID')),owner_external_id=text(r.get('提供主体')) or '',category=text(r.get('资源类别')),description=text(r.get('资源说明')),region=text(r.get('地域')),applicable_to=text(r.get('适用对象')),visibility=text(r.get('权限')) or '内部',verification_status=text(r.get('核验状态')) or '待核验'))
        summary['历史事件种子']=import_sheet(db,wb['历史事件种子'],'历史事件种子',HistoricalEvent,lambda r:HistoricalEvent(external_id=text(r.get('事件ID')),event_date=text(r.get('日期')),name=text(r.get('事件名称')) or '未命名事件',event_type=text(r.get('事件类型')),related_entity=text(r.get('关联主体')),fact_summary=text(r.get('事实摘要')),system_use=text(r.get('系统用途')),visibility=text(r.get('权限')) or '内部',verification_status=text(r.get('核验状态')) or '待核验'))
        summary['项目池种子']=import_sheet(db,wb['项目池种子'],'项目池种子',ProjectPool,lambda r:ProjectPool(external_id=text(r.get('项目池ID')),name=text(r.get('项目池名称')) or '未命名项目池',project_type=text(r.get('类型')),owner_external_id=text(r.get('归属主体')),focus_tags=text(r.get('重点标签')),typical_needs=text(r.get('典型需求')),target_actions=text(r.get('目标动作')),visibility=text(r.get('权限')) or '内部',status=text(r.get('状态')) or '待整理'))
        summary['关系种子库']=import_sheet(db,wb['关系种子库'],'关系种子库',Relation,lambda r:Relation(external_id=text(r.get('关系ID')),source_external_id=text(r.get('起点')) or '',relation_type=text(r.get('关系类型')) or '未定义关系',target_external_id=text(r.get('终点')) or '',period=text(r.get('时间')),evidence_source=text(r.get('证据来源')),visibility=text(r.get('权限')) or '内部',verification_status=text(r.get('核验状态')) or '待核验'))
        summary['校准行动清单']=import_sheet(db,wb['校准行动清单'],'校准行动清单',ActionItem,lambda r:ActionItem(external_id=text(r.get('行动ID')),task=text(r.get('任务')) or '未命名任务',target_external_id=text(r.get('对象')),completion_standard=text(r.get('完成标准')),owner=text(r.get('负责人')),priority=text(r.get('优先级')),status=text(r.get('状态')) or '未开始',suggested_deadline=text(r.get('建议截止'))))
        db.add(ImportLog(filename=fp.name,result_json=json.dumps(summary,ensure_ascii=False))); db.commit()
    finally:
        db.close()
    print('\nIMPORT COMPLETE'); print('='*72)
    for sheet,result in summary.items():
        print(f"{sheet}: 新增 {result['added']}，跳过 {result['skipped']}，失败 {result['failed']}")
        for e in result['errors']: print(' -',e)
    print('='*72)
if __name__=='__main__': main()
