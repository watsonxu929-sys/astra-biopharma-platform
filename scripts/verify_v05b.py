from __future__ import annotations

import io
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)
    print(f"[PASS] {message}")


def build_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    wb = Workbook()
    ws = wb.active
    ws.title = "会员"
    ws.append(["姓名", "机构", "职位", "手机", "邮箱", "可提供资源", "合作需求"])
    ws.append(["张测试", "测试生物", "BD负责人", "13800000001", "zhang@example.com", "投资机构资源", "融资需求"])
    image_bytes = io.BytesIO()
    Image.new("RGB", (120, 120), "white").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    pic = XLImage(image_bytes)
    pic.anchor = "H2"
    ws.add_image(pic)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_docx() -> bytes:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    headers = ["姓名", "机构", "职位", "手机", "邮箱", "可提供资源", "合作需求"]
    values = ["李测试", "测试医药", "研发负责人", "13900000002", "li@example.com", "技术平台", "临床资源"]
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for idx, value in enumerate(values):
        table.rows[1].cells[idx].text = value
    image_bytes = io.BytesIO()
    Image.new("RGB", (100, 100), "gray").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    doc.add_picture(image_bytes, width=Inches(1))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def main() -> int:
    original_cwd = Path.cwd()
    with tempfile.TemporaryDirectory(prefix="v05b_verify_", ignore_cleanup_errors=True) as tmp:
        work = Path(tmp)
        (work / "data").mkdir(parents=True, exist_ok=True)
        os.chdir(work)
        os.environ["APP_DB_PATH"] = str(work / "data" / "app.db")
        os.environ["APP_AUTH_DISABLED"] = "1"

        from app.services.member_import_service import parse_text, parse_uploaded_document

        pasted = "\n".join([
            "\t".join(["姓名", "机构", "职位", "手机", "邮箱", "可提供资源", "合作需求"]),
            "\t".join(["王一", "甲生物", "创始人", "13700000001", "w1@example.com", "产业资源", "融资"]),
            "\t".join(["王二", "乙医药", "投资负责人", "13600000002", "w2@example.com", "投资资源", "项目合作"]),
        ])
        drafts = parse_text(pasted)
        check(len(drafts) == 2 and drafts[0]["name"] == "王一", "粘贴表格可拆分为独立会员草稿")

        xlsx_path = work / "members.xlsx"
        xlsx_path.write_bytes(build_xlsx())
        xlsx_drafts, xlsx_images = parse_uploaded_document(xlsx_path)
        check(len(xlsx_drafts) == 1 and xlsx_drafts[0]["organization_name"] == "测试生物", "Excel 表头自动映射")
        check(len(xlsx_images) == 1, "Excel 内嵌图片可提取")

        docx_path = work / "members.docx"
        docx_path.write_bytes(build_docx())
        docx_drafts, docx_images = parse_uploaded_document(docx_path)
        check(any(item["name"] == "李测试" for item in docx_drafts), "Word 表格可解析")
        check(len(docx_images) == 1, "Word 内嵌图片可提取")

        from fastapi.testclient import TestClient
        import scripts.apply_v04b as apply_v04b_module
        apply_v04b_module.apply_v04b.__defaults__ = (work / "data" / "app.db",)
        from app.main import app

        client = TestClient(app)
        check(client.get("/v05b/health").status_code == 200, "v0.5B 健康检查可访问")
        check(client.get("/club/import").status_code == 200, "会员智能导入页面可访问")
        template = client.get("/club/import/template.xlsx")
        check(template.status_code == 200 and len(template.content) > 1000, "Excel 标准模板可下载")

        response = client.post(
            "/club/import/text",
            data={
                "raw_text": "姓名：赵验证\n机构：验证生物\n职位：CEO\n手机：13500000003\n邮箱：zhao@example.com\n可提供资源：产业资源\n合作需求：融资合作",
                "source_url": "",
            },
            follow_redirects=False,
        )
        check(response.status_code == 303 and "/club/import/jobs/" in response.headers["location"], "粘贴资料可创建导入任务")
        job_url = response.headers["location"]
        check(client.get(job_url).status_code == 200, "导入核对工作台可访问")

        db_path = work / "data" / "app.db"
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        draft = conn.execute("SELECT * FROM v05b_import_drafts ORDER BY id DESC LIMIT 1").fetchone()
        job = conn.execute("SELECT * FROM v05b_import_jobs ORDER BY id DESC LIMIT 1").fetchone()
        conn.close()
        check(bool(draft and job), "导入任务与草稿已持久化")

        form = {f"selected_{draft['id']}": "1", f"duplicate_action_{draft['id']}": "create", f"existing_person_id_{draft['id']}": "0", f"existing_org_id_{draft['id']}": "0"}
        for field in [
            "name", "organization_name", "title", "mobile", "email", "wechat", "city",
            "industry_tags", "expertise_tags", "offered_resources", "cooperation_needs",
            "self_introduction", "referral_source", "referrer_name", "preferred_contact_method",
            "member_level", "member_status", "owner",
        ]:
            form[f"{field}_{draft['id']}"] = draft[field] or ""
        save = client.post(f"/club/import/jobs/{job['id']}/save", data=form, follow_redirects=False)
        check(save.status_code == 303, "已核对草稿可批量保存")

        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        membership = conn.execute(
            """
            SELECT m.*,p.name,o.standard_name,c.mobile,c.email
            FROM v04f_club_memberships m
            JOIN people p ON p.id=m.person_id
            LEFT JOIN organizations o ON o.id=m.organization_id
            LEFT JOIN v05b_member_contacts c ON c.membership_id=m.id
            ORDER BY m.id DESC LIMIT 1
            """
        ).fetchone()
        needs = conn.execute("SELECT COUNT(*) FROM v04f_club_needs WHERE membership_id=?", (membership["id"],)).fetchone()[0]
        offers = conn.execute("SELECT COUNT(*) FROM v04f_club_offerings WHERE membership_id=?", (membership["id"],)).fetchone()[0]
        conn.close()
        check(membership["name"] == "赵验证" and membership["standard_name"] == "验证生物", "人物、机构与会员身份联动保存")
        check(membership["mobile"] == "13500000003" and membership["email"] == "zhao@example.com", "会员联系方式进入隐私表")
        check(needs == 1 and offers == 1, "会员需求与可提供资源同步保存")

        from PIL import Image
        avatar = io.BytesIO()
        Image.new("RGB", (200, 200), "blue").save(avatar, format="PNG")
        avatar_upload = client.post(
            f"/club/members/{membership['id']}/avatar",
            files={"file": ("avatar.png", avatar.getvalue(), "image/png")},
            follow_redirects=False,
        )
        check(avatar_upload.status_code == 303, "会员头像可上传")
        conn = sqlite3.connect(db_path)
        asset = conn.execute("SELECT id FROM v05b_media_assets WHERE membership_id=? AND is_active=1", (membership["id"],)).fetchone()
        conn.close()
        check(bool(asset), "头像元数据已保存")
        check(client.get(f"/club/media/{asset[0]}?thumb=1").status_code == 200, "头像缩略图可读取")

        duplicate = client.post(
            "/club/import/text",
            data={"raw_text": "姓名：赵验证\n机构：验证生物\n手机：13500000003\n邮箱：zhao@example.com"},
            follow_redirects=False,
        )
        duplicate_job = duplicate.headers["location"]
        page = client.get(duplicate_job)
        check(page.status_code == 200 and "检测到现有会员" in page.text, "手机号或邮箱重复可被识别并提示")

        check(client.get(f"/club/members/{membership['id']}").status_code == 200, "会员档案兼容头像和隐私联系方式")

    os.chdir(original_cwd)
    print("v0.5B verification passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
